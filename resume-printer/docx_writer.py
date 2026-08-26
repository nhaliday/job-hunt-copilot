"""Write a resume DOCX from the intermediate HTML, mirroring resume.css.

The DOCX is generated programmatically (python-docx, with raw OOXML for what
it can't express: rule borders, custom bullet numbering, hyperlinks) rather
than via pandoc --reference-doc — named-style mapping can't reproduce the
per-element formatting this design needs (centered ruled headings, borderless
two-column entry tables).

Input is the same post-filter HTML that fit.py renders, so PDF and DOCX
content always agree. All type sizes derive from one base font size via the
em ratios in resume.css; absolute spacing (margins, rule padding) stays fixed,
also mirroring the CSS.

Page fit: start from the PDF's fitted size minus a safety buffer (Word's line
breaking differs slightly from WeasyPrint's), then, when LibreOffice is
available, verify the page count via headless DOCX->PDF conversion and step
the size down until the target page count is met or --min-pt is reached.
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor, Twips
from lxml import html as lhtml

INK = "171717"
FONT = "EB Garamond"
LINE_HEIGHT = 1.29
EM_CONTACT = 0.895
EM_NAME = 2.368
EM_H2 = 1.158
BULLET_INDENT_TWIPS = 280  # ~ ul padding-left: 14pt
# \u25c6 U+25C6, same as resume.css: it's the diamond EB Garamond actually covers
# (\u2666 U+2666 isn't in the font, and its emoji presentation renders red)
BULLET = "\u25c6"


def _half_pt(pt):
    return Pt(round(pt * 2) / 2)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _ppr_insert(par, child, successors):
    """Insert child into the paragraph's pPr at its schema position."""
    ppr = par._p.get_or_add_pPr()
    for tag in successors:
        succ = ppr.find(qn(tag))
        if succ is not None:
            succ.addprevious(child)
            return
    ppr.append(child)


def _rule(par, edge, space_pt):
    """Double-line paragraph border, the section-separator rule."""
    pbdr = _el("w:pBdr")
    pbdr.append(_el(f"w:{edge}", val="double", sz=1, space=space_pt, color=INK))
    _ppr_insert(par, pbdr, ("w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr"))


def _letter_space(run, pt):
    run._r.get_or_add_rPr().append(_el("w:spacing", val=int(pt * 20)))


def _font_fallback(doc, name, alt):
    """Declare a substitute in fontTable.xml (w:altName): machines without
    `name` installed fall back to `alt` instead of Word's default serif —
    Garamond ships with Office and is close enough in metrics that the
    one-page fit usually survives."""
    from lxml import etree

    part = doc.part.part_related_by(RT.FONT_TABLE)
    root = etree.fromstring(part.blob)
    font = etree.SubElement(root, qn("w:font"))
    font.set(qn("w:name"), name)
    for tag, val in (
        ("w:altName", alt),
        ("w:family", "roman"),
        ("w:pitch", "variable"),
    ):
        e = etree.SubElement(font, qn(tag))
        e.set(qn("w:val"), val)
    part._blob = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def _fmt(par, *, line_pt, align=None, before=0.0, after=0.0):
    """line_pt is an exact line height in points (CSS line-height semantics —
    Word's 'multiple' spacing would compound with the font's built-in leading
    and come out ~17% taller than the PDF)."""
    pf = par.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = Pt(line_pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    return par


# --- inline content ------------------------------------------------------


def _segments(el):
    """Flatten an element to [text, bold, italic, href] segments."""
    segs = []

    def walk(node, bold, italic, href):
        if node.text:
            segs.append([node.text, bold, italic, href])
        for c in node:
            if not isinstance(c.tag, str):
                pass
            elif c.tag in ("strong", "b"):
                walk(c, True, italic, href)
            elif c.tag in ("em", "i"):
                walk(c, bold, True, href)
            elif c.tag == "a":
                walk(c, bold, italic, c.get("href"))
            else:
                if c.tag in ("p", "br") and segs:
                    segs.append([" ", bold, italic, None])
                walk(c, bold, italic, href)
            if c.tail:
                segs.append([c.tail, bold, italic, href])

    walk(el, False, False, None)
    for s in segs:
        s[0] = re.sub(r"\s+", " ", s[0])
    while segs and not segs[0][0].strip():
        segs.pop(0)
    while segs and not segs[-1][0].strip():
        segs.pop()
    if segs:
        segs[0][0] = segs[0][0].lstrip()
        segs[-1][0] = segs[-1][0].rstrip()
    return segs


def _run(par, text, size_pt, bold, italic, caps):
    r = par.add_run(text)
    r.font.size = _half_pt(size_pt)
    if bold:
        r.font.bold = True
    if italic:
        r.font.italic = True
    if caps:
        r.font.all_caps = True
    return r


def _emit(par, segs, size_pt, *, bold=False, italic=False, caps=False, spacing_pt=0):
    for text, b, i, href in segs:
        if href:
            r_id = par.part.relate_to(href, RT.HYPERLINK, is_external=True)
            link = _el("w:hyperlink")
            link.set(qn("r:id"), r_id)
            par._p.append(link)
            r = _run(par, text, size_pt, bold or b, italic or i, caps)
            link.append(r._r)
        else:
            r = _run(par, text, size_pt, bold or b, italic or i, caps)
        if spacing_pt:
            _letter_space(r, spacing_pt)


# --- bullets -------------------------------------------------------------


def _bullet_num_id(doc):
    """Register a bullet list definition; None if the package has no
    numbering part (fall back to literal markers)."""
    try:
        numbering = doc.part.numbering_part.element
    except KeyError, NotImplementedError:
        return None
    abses = numbering.findall(qn("w:abstractNum"))
    nums = numbering.findall(qn("w:num"))
    abs_id = max((int(a.get(qn("w:abstractNumId"))) for a in abses), default=-1) + 1
    num_id = max((int(n.get(qn("w:numId"))) for n in nums), default=0) + 1

    lvl = _el("w:lvl", ilvl=0)
    lvl.append(_el("w:start", val=1))
    lvl.append(_el("w:numFmt", val="bullet"))
    lvl.append(_el("w:lvlText", val=BULLET))
    lvl.append(_el("w:lvlJc", val="left"))
    ppr = _el("w:pPr")
    ppr.append(_el("w:ind", left=BULLET_INDENT_TWIPS, hanging=BULLET_INDENT_TWIPS))
    lvl.append(ppr)
    # pin the marker's font and color — an unstyled ♦ falls back to the
    # emoji font (red diamond) in both Word and LibreOffice
    rpr = _el("w:rPr")
    fonts = _el("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    rpr.append(fonts)
    rpr.append(_el("w:color", val=INK))
    lvl.append(rpr)
    absn = _el("w:abstractNum", abstractNumId=abs_id)
    absn.append(lvl)
    if nums:
        nums[0].addprevious(absn)
    else:
        numbering.append(absn)
    num = _el("w:num", numId=num_id)
    num.append(_el("w:abstractNumId", val=abs_id))
    numbering.append(num)
    return num_id


def _bullet_para(doc, li, size_pt, num_id, *, before, after):
    par = _fmt(
        doc.add_paragraph(), line_pt=LINE_HEIGHT * size_pt, before=before, after=after
    )
    pf = par.paragraph_format
    pf.left_indent = Twips(BULLET_INDENT_TWIPS)
    if num_id is not None:
        npr = _el("w:numPr")
        npr.append(_el("w:ilvl", val=0))
        npr.append(_el("w:numId", val=num_id))
        _ppr_insert(
            par,
            npr,
            ("w:pBdr", "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr"),
        )
    else:
        pf.first_line_indent = Twips(-BULLET_INDENT_TWIPS)
        pf.tab_stops.add_tab_stop(Twips(BULLET_INDENT_TWIPS))
        _run(par, BULLET + "\t", size_pt, False, False, False)
    _emit(par, _segments(li), size_pt)


# --- entry tables --------------------------------------------------------


def _two_col_table(
    doc, sec, left_td, right_td, size_pt, *, caps, italic, before, after
):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False  # fixed layout: honor the explicit column widths
    tblpr = tbl._tbl.tblPr

    def insert(el):
        look = tblpr.find(qn("w:tblLook"))
        if look is not None:
            look.addprevious(el)
        else:
            tblpr.append(el)

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = _el("w:tblW")
        insert(tblw)
    tblw.set(qn("w:type"), "pct")
    tblw.set(qn("w:w"), "5000")
    borders = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(f"w:{edge}", val="none", sz=0, color="auto"))
    insert(borders)
    margins = _el("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        margins.append(_el(f"w:{edge}", w=0, type="dxa"))
    insert(margins)

    # The CSS auto-sizes these columns with a nowrap right cell; Word tables
    # need explicit widths, so give the right column its estimated text width
    # (~0.55em average glyph advance for bold EB Garamond), clamped.
    content = int(sec.page_width - sec.left_margin - sec.right_margin)
    right_text = right_td.text_content().strip()
    needed = int(Pt(len(right_text) * 0.55 * size_pt)) + int(Pt(6))
    right_w = min(max(needed, content // 6), content // 2)
    left, right = tbl.rows[0].cells
    tbl.columns[0].width = left.width = Emu(content - right_w)
    tbl.columns[1].width = right.width = Emu(right_w)
    lp = _fmt(
        left.paragraphs[0], line_pt=LINE_HEIGHT * size_pt, before=before, after=after
    )
    _emit(lp, _segments(left_td), size_pt, bold=True, italic=italic, caps=caps)
    rp = _fmt(
        right.paragraphs[0],
        line_pt=LINE_HEIGHT * size_pt,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        before=before,
        after=after,
    )
    _emit(rp, _segments(right_td), size_pt, bold=True, italic=italic)


# --- document ------------------------------------------------------------


def write_docx(html_path, out_path, base_pt):
    root = lhtml.parse(html_path).getroot()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)  # A4, as resume.css @page
    sec.top_margin = sec.bottom_margin = Mm(11)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.header_distance = sec.footer_distance = Emu(0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = _half_pt(base_pt)
    normal.font.color.rgb = RGBColor(0x17, 0x17, 0x17)
    normal.paragraph_format.line_spacing = Pt(LINE_HEIGHT * base_pt)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    num_id = _bullet_num_id(doc)
    center = WD_ALIGN_PARAGRAPH.CENTER

    header = root.body.find("header")
    name_text = ""
    if header is not None:
        first = None
        contact = header.find_class("header-contact")
        if contact:
            par = _fmt(
                doc.add_paragraph(),
                line_pt=1.4 * base_pt * EM_CONTACT,
                align=center,
                after=2,
            )
            _emit(par, _segments(contact[0]), base_pt * EM_CONTACT)
            first = par
        name = header.find_class("header-name")
        if name:
            name_text = name[0].text_content().strip()
            par = _fmt(
                doc.add_paragraph(), line_pt=1.15 * base_pt * EM_NAME, align=center
            )
            _emit(
                par,
                _segments(name[0]),
                base_pt * EM_NAME,
                bold=True,
                caps=True,
                spacing_pt=0.5,
            )
            first = first or par
        subtitle = header.find_class("header-subtitle")
        if subtitle:
            par = _fmt(
                doc.add_paragraph(),
                line_pt=LINE_HEIGHT * base_pt,
                align=center,
                after=6,
            )
            _emit(par, _segments(subtitle[0]), base_pt, italic=True)
            first = first or par
        if first is not None:
            _rule(first, "top", 6)  # .resume-header border-top + padding-top

    for el in root.body.iterchildren():
        if el is header or not isinstance(el.tag, str):
            continue
        if el.tag == "h2":
            par = _fmt(
                doc.add_paragraph(),
                line_pt=LINE_HEIGHT * base_pt * EM_H2,
                align=center,
                before=12,
                after=6,
            )
            _emit(
                par,
                _segments(el),
                base_pt * EM_H2,
                bold=True,
                caps=True,
                spacing_pt=0.5,
            )
            _rule(par, "bottom", 4)  # h2 border-bottom + padding-bottom
        elif el.tag == "table":
            tds = el.xpath(".//td")
            if len(tds) != 2:
                continue
            if "entry-org" in el.get("class", ""):
                _two_col_table(
                    doc, sec, *tds, base_pt, caps=False, italic=True, before=0, after=3
                )
            else:
                _two_col_table(
                    doc, sec, *tds, base_pt, caps=True, italic=False, before=2, after=0
                )
        elif el.tag in ("ul", "ol"):
            lis = el.findall("li")
            for i, li in enumerate(lis):
                _bullet_para(
                    doc,
                    li,
                    base_pt,
                    num_id,
                    before=1 if i == 0 else 0,
                    after=2 if i == len(lis) - 1 else 0.5,
                )
        elif el.tag == "h3":
            par = _fmt(doc.add_paragraph(), line_pt=LINE_HEIGHT * base_pt, before=2)
            _emit(par, _segments(el), base_pt, bold=True, caps=True)
        elif el.tag == "p":
            par = _fmt(
                doc.add_paragraph(),
                line_pt=LINE_HEIGHT * base_pt,
                before=base_pt / 2,
                after=base_pt / 2,
            )
            _emit(par, _segments(el), base_pt)

    _font_fallback(doc, FONT, "Garamond")
    doc.core_properties.author = name_text
    doc.core_properties.title = name_text
    doc.save(out_path)


# --- page fit ------------------------------------------------------------


def find_soffice():
    path = shutil.which("soffice")
    if path:
        return path
    mac_app = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return mac_app if os.path.exists(mac_app) else None


def count_pages(soffice, docx_path):
    """Headless LibreOffice DOCX->PDF, return the page count. A per-call user
    profile keeps parallel build jobs from fighting over the default one."""
    import pdfplumber

    with tempfile.TemporaryDirectory() as td:
        profile = Path(td, "profile").as_uri()
        subprocess.run(
            [
                soffice,
                "--headless",
                f"-env:UserInstallation={profile}",
                "--convert-to",
                "pdf",
                "--outdir",
                td,
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
        pdf = Path(td, Path(docx_path).stem + ".pdf")
        if not pdf.exists():
            raise RuntimeError(f"soffice produced no PDF for {docx_path}")
        with pdfplumber.open(pdf) as converted:
            return len(converted.pages)


def main():
    parser = argparse.ArgumentParser(
        description="Generate a resume DOCX from the intermediate HTML"
    )
    parser.add_argument("html", help="Intermediate HTML (post filter/template)")
    parser.add_argument("docx", help="Output DOCX file")
    parser.add_argument(
        "--pdf-fit-pt",
        type=float,
        default=None,
        help="Font size fit.py chose for the PDF (warm start)",
    )
    parser.add_argument("--pages", type=int, default=1, help="Target page count")
    parser.add_argument("--min-pt", type=float, default=10.0)
    parser.add_argument("--max-pt", type=float, default=12.0)
    parser.add_argument(
        "--buffer-pt",
        type=float,
        default=0.3,
        help="Safety margin under the PDF fit (Word breaks lines differently)",
    )
    parser.add_argument("--step-pt", type=float, default=0.25)
    args = parser.parse_args()

    if args.pdf_fit_pt is not None:
        size = min(args.max_pt, max(args.min_pt, args.pdf_fit_pt - args.buffer_pt))
    else:
        size = args.min_pt
    write_docx(args.html, args.docx, size)

    soffice = find_soffice()
    if soffice is None:
        print(
            f"{size:.2f}pt (docx; LibreOffice not found — page fit unverified)",
            file=sys.stderr,
        )
        return
    try:
        pages = count_pages(soffice, args.docx)
        while pages > args.pages and size > args.min_pt:
            size = max(args.min_pt, size - args.step_pt)
            write_docx(args.html, args.docx, size)
            pages = count_pages(soffice, args.docx)
    except (subprocess.SubprocessError, RuntimeError) as exc:
        print(f"    WARN: docx page-fit check failed: {exc}", file=sys.stderr)
        return
    if pages > args.pages:
        print(
            f"    WARN: docx still {pages} pages at minimum {args.min_pt}pt",
            file=sys.stderr,
        )
    print(f"{size:.2f}pt (docx, {pages}p via soffice)", file=sys.stderr)


if __name__ == "__main__":
    main()
